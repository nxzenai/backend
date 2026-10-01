"""Downloadable views of completed native lab batch predictions."""
from __future__ import annotations

import csv
import io
import re
from collections import Counter
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook


def _mapped(value: Any, mapping: dict[str, Any]) -> Any:
    if value is None:
        return None
    label = mapping.get(str(value))
    if label is not None and str(label).strip():
        return str(label)
    return str(value) if not str(value).strip().replace(".", "", 1).isdigit() else None


def _probability(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if 0 <= number <= 1 else None


def _percent(value: Any) -> float | None:
    probability = _probability(value)
    return round(probability * 100, 2) if probability is not None else None


def _column(label: Any) -> str:
    return "prob_" + (re.sub(r"[^a-z0-9]+", "_", str(label).lower()).strip("_") or "class") + "_pct"


def _reference(value: Any) -> Any:
    return re.split(r"[/\\]", str(value))[-1] if value is not None else None


def _source_rows(contents: bytes) -> list[dict[str, Any]]:
    # Keep source values as entered; only appended result columns are rounded.
    return list(csv.DictReader(io.StringIO(contents.decode("utf-8-sig"), newline="")))


def build_prediction_export(module: str, native: dict[str, Any], source: bytes | None) -> tuple[bytes, bytes, dict[str, Any]]:
    task = str(native.get("task") or ("classification" if module == "autonlp" else "image_classification" if module == "autodl" else ""))
    if module in {"automl", "autonlp"}:
        if source is None:
            raise ValueError("Prediction source file is unavailable.")
        records = _source_rows(source)
    else:
        records = []
    output: list[dict[str, Any]] = []
    labels: list[str] = []
    confidences: list[float] = []
    mapping: dict[str, Any] = {}
    if module == "automl":
        target = native.get("target_metadata") or {}
        mapping = next((target[key] for key in ("class_meanings", "label_mapping", "value_meanings", "display_mapping") if isinstance(target.get(key), dict)), {})
        predictions = native.get("predictions") or []
        classes = native.get("classes") or []
        if len(records) != len(predictions):
            raise ValueError("Native prediction count does not match the test data.")
        probabilities = native.get("probabilities") or []
        class_names = [_mapped(item, mapping) for item in classes]
        probability_columns = [_column(item) for item in class_names] if len(classes) <= 10 and all(class_names) else []
        if len(set(probability_columns)) != len(probability_columns):
            probability_columns = []
        for index, (record, prediction) in enumerate(zip(records, predictions)):
            row = dict(record)
            isolated = (native.get("row_results") or [None] * len(records))[index]
            error = isolated.get("error") if isinstance(isolated, dict) else None
            if task == "classification":
                label = _mapped(prediction, mapping) if not error else None
                confidence = _probability((native.get("prediction_confidences") or [None] * len(records))[index])
                row.update(predicted_class=prediction, predicted_label=label, prediction_confidence_pct=_percent(confidence))
                if not error and len(probability_columns) == len(classes) and index < len(probabilities):
                    scores = probabilities[index]
                    if isinstance(scores, list) and len(scores) == len(classes):
                        row.update({name: _percent(score) for name, score in zip(probability_columns, scores)})
                elif not error and len(classes) > 10 and index < len(probabilities):
                    scores = probabilities[index]
                    if isinstance(scores, list) and len(scores) == len(classes):
                        ranked = sorted(range(len(classes)), key=lambda position: scores[position], reverse=True)
                        if len(ranked) > 1:
                            row["second_best_label"] = class_names[ranked[1]]
                            row["second_best_confidence_pct"] = _percent(scores[ranked[1]])
                if label is not None or prediction is not None:
                    labels.append(str(label if label is not None else prediction))
                if confidence is not None:
                    confidences.append(confidence)
            elif task == "regression":
                row["predicted_value"] = round(float(prediction), 2) if prediction is not None else None
            elif task == "clustering":
                segment = (native.get("cluster_name_mapping") or {}).get(str(prediction)) if not error else None
                row.update(cluster_id=prediction, cluster_name=segment)
                if segment is not None:
                    labels.append(str(segment))
            row.update(prediction_status="failed" if error else "success", prediction_error=error)
            output.append(row)
    elif module == "autonlp":
        mapping = native.get("label_display_mapping") or {}
        by_index = {int(item["row_index"]): item for item in native.get("rows") or []}
        probability_columns = {_column(label): str(label) for label in mapping.values()} if 1 < len(mapping) <= 3 else {}
        if len(probability_columns) != len(mapping):
            probability_columns = {}
        for index, record in enumerate(records):
            item = by_index.get(index) or {}
            row = dict(record)
            error = item.get("error") or ("Native result is missing for this row." if not item else None)
            technical = item.get("technical_label")
            label = _mapped(technical, mapping) or _mapped(item.get("predicted_label"), mapping)
            confidence = _probability(item.get("model_score")) if not error else None
            row.update(predicted_class=technical, predicted_label=label if not error else None,
                       prediction_confidence_pct=_percent(confidence), prediction_status="failed" if error else "success",
                       prediction_error=error)
            scores = item.get("probabilities") or []
            if not error and probability_columns and len(scores) == len(mapping):
                available = {str(score.get("label")): _probability(score.get("probability")) for score in scores}
                if set(probability_columns.values()) == set(available):
                    row.update({column: _percent(available[name]) for column, name in probability_columns.items()})
            elif not error and len(mapping) > 10 and len(scores) > 1:
                row["second_best_label"] = scores[1].get("label")
                row["second_best_confidence_pct"] = _percent(scores[1].get("probability"))
            if not error and (label or technical is not None):
                labels.append(str(label if label else technical))
            if confidence is not None:
                confidences.append(confidence)
            output.append(row)
    else:
        classes = native.get("class_labels") or []
        mapping = {str(index): str(label) for index, label in enumerate(classes)}
        combined = [*native.get("predictions", []), *native.get("errors", [])]
        for item in sorted(combined, key=lambda value: value.get("row", 0)):
            raw_class = item.get("predicted_class", item.get("predicted_category"))
            label = str(raw_class) if raw_class is not None and str(raw_class) in [str(value) for value in classes] else None
            if label is None and raw_class is not None and str(raw_class).isdigit() and int(raw_class) < len(classes):
                label = str(classes[int(raw_class)])
            error = item.get("message") or item.get("error")
            confidence = _probability(item.get("confidence")) if not error else None
            output.append({"image_name": item.get("image_name"), "predicted_class": raw_class,
                           "predicted_label": label, "prediction_confidence_pct": _percent(confidence),
                           "prediction_status": "failed" if error else "success", "prediction_error": error})
            if not error and (label or raw_class is not None):
                labels.append(str(label if label else raw_class))
            if confidence is not None:
                confidences.append(confidence)
    if not output:
        raise ValueError("Native batch prediction returned no rows.")
    columns = list(dict.fromkeys(key for row in output for key in row))
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=columns)
    writer.writeheader()
    writer.writerows(output)
    workbook = Workbook()
    predictions_sheet = workbook.active
    predictions_sheet.title = "Predictions"
    predictions_sheet.append(columns)
    for row in output:
        predictions_sheet.append([row.get(key) for key in columns])
    successes = sum(row["prediction_status"] == "success" for row in output)
    summary = workbook.create_sheet("Summary")
    summary.append(["Field", "Value"])
    for key, value in (("module", module), ("task", task), ("model", _reference(native.get("model_name") or native.get("model_filename") or native.get("model_id") or native.get("run_id"))),
                       ("rows processed", len(output)), ("successful predictions", successes), ("failed predictions", len(output) - successes),
                       ("average confidence (%)", round(100 * sum(confidences) / len(confidences), 2) if confidences else None)):
        summary.append([key, value])
    for label, count in Counter(labels).items():
        summary.append([f"class: {label}", f"{count} ({count / successes:.2%})" if successes else count])
    info = workbook.create_sheet("Model_Info")
    info.append(["Field", "Value"])
    for key, value in (("model/artifact/run reference", _reference(native.get("model_filename") or native.get("model_id") or native.get("run_id"))),
                       ("task type", task), ("target", (native.get("target_metadata") or {}).get("name") if module == "automl" else native.get("text_column")),
                       ("feature schema", ", ".join(native.get("input_features") or [])),
                       ("label mapping", ", ".join(f"{key}: {value}" for key, value in mapping.items())),
                       ("prediction timestamp", datetime.now(UTC).isoformat())):
        info.append([key, value])
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    stats = {"rows": len(output), "successes": successes, "failures": len(output) - successes,
             "distribution": dict(Counter(labels)), "average_confidence": sum(confidences) / len(confidences) if confidences else None}
    return csv_buffer.getvalue().encode("utf-8-sig"), xlsx_buffer.getvalue(), stats
